import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

# Sales data
data = {
    "Product": ["Samsung Galaxy S24", "HP Laptop", "boAt Earbuds", "LG Smart TV"],
    "Category": ["Mobile", "Laptop", "Accessories", "Television"],
    "Quantity Sold": [2, 3, 10, 1],
    "Selling Price": [55000, 60000, 2000, 45000]
}

df = pd.DataFrame(data)
df["Total Revenue"] = df["Quantity Sold"] * df["Selling Price"]

# Category summary
summary = df.groupby("Category", as_index=False)["Total Revenue"].sum()

# Export both sheets
file_path = "data/processed/SmartBiz_Professional_Report.xlsx"

with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
    df.to_excel(writer, sheet_name="Sales Data", index=False)
    summary.to_excel(writer, sheet_name="Category Summary", index=False)

# Apply Excel formatting
workbook = load_workbook(file_path)

for sheet in workbook.worksheets:
    sheet.freeze_panes = "A2"

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            start_color="1F4E78",
            end_color="1F4E78",
            fill_type="solid"
        )
        cell.alignment = Alignment(horizontal="center")

    for column in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = max_length + 4

workbook.save(file_path)

print("Professional Excel report created successfully!")
print("File:", file_path)
