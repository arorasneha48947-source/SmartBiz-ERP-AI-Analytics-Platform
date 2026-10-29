import pandas as pd
import matplotlib.pyplot as plt
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parents[1]
INPUT_FILE = BASE_DIR / "data/processed/SmartBiz_Business_KPI_Dashboard.xlsx"
OUTPUT_FILE = BASE_DIR / "data/processed/SmartBiz_Professional_Dashboard.xlsx"

# Read the existing KPI report
df = pd.read_excel(INPUT_FILE)

# Create workbook
with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
    df.to_excel(writer, sheet_name="Business Dashboard", index=False)

# Open workbook for formatting
wb = load_workbook(OUTPUT_FILE)
ws = wb["Business Dashboard"]

# Title
ws.insert_rows(1, 2)
ws.merge_cells("A1:B1")
ws["A1"] = "SMARTBIZ ELECTRONICS"
ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
ws["A1"].fill = PatternFill("solid", fgColor="172554")
ws["A1"].alignment = Alignment(horizontal="center")

ws.merge_cells("A2:B2")
ws["A2"] = "Business KPI Dashboard | Practice Project"
ws["A2"].font = Font(size=11, italic=True, color="FFFFFF")
ws["A2"].fill = PatternFill("solid", fgColor="1E40AF")
ws["A2"].alignment = Alignment(horizontal="center")

# Header formatting
header_fill = PatternFill("solid", fgColor="2563EB")
for cell in ws[3]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")

# KPI row formatting
for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
    for cell in row:
        cell.font = Font(size=11, color="172554")
        cell.alignment = Alignment(vertical="center")
        cell.border = Border(bottom=Side(style="thin", color="D1D5DB"))

    row[1].font = Font(size=11, bold=True, color="166534")

# Column widths
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 22

# Add a chart using the KPI values
chart = BarChart()
chart.type = "bar"
chart.style = 10
chart.title = "Business KPI Overview"
chart.y_axis.title = "Business KPIs"
chart.x_axis.title = "Values"

data = Reference(ws, min_col=2, min_row=3, max_row=ws.max_row)
categories = Reference(ws, min_col=1, min_row=4, max_row=ws.max_row)
chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)
chart.height = 10
chart.width = 18

ws.add_chart(chart, "D4")

ws.freeze_panes = "A4"
ws.sheet_view.showGridLines = False

# Save final workbook
wb.save(OUTPUT_FILE)

print("Professional dashboard created successfully!")
print(f"Saved at: {OUTPUT_FILE}")
